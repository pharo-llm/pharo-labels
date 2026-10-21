"""Expert workbooks (Excel): writing them and importing them back without losing who decided what.

Every workbook has a ``README`` sheet with the instructions and a ``reviewer_id`` cell the expert fills in
(the file name is the fallback: ``<reviewer>.xlsx``). Input columns have a yellow header and drop-down
lists where the answer is a closed choice; read-only columns are grey. Rows carry a stable ``row_key``
that must not be edited: on import, unknown or duplicated keys, missing reviewer identity and invalid
choices are reported (``import_report.md``), never silently fixed. The imported rows keep the reviewer,
the file name, its sha256 and the import time.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .common import file_sha256, now_iso

INPUT_FILL = PatternFill("solid", fgColor="FFF4C2")
READ_FILL = PatternFill("solid", fgColor="ECEBE7")
HEAD_FONT = Font(bold=True)


def new_workbook(title: str, instructions: list[str], extra: dict | None = None) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "README"
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14)
    ws["A3"] = "reviewer_id"
    ws["A3"].font = HEAD_FONT
    ws["B3"].fill = INPUT_FILL
    ws["C3"] = "← your name or initials (required)"
    r = 5
    for line in instructions:
        ws.cell(row=r, column=1, value=line).alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
        ws.row_dimensions[r].height = 15 * max(1, len(line) // 110 + 1)
        r += 1
    for k, v in (extra or {}).items():
        r += 1
        ws.cell(row=r, column=1, value=k).font = HEAD_FONT
        ws.cell(row=r, column=2, value=str(v))
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 30
    return wb


def add_sheet(wb: Workbook, name: str, rows: list[dict], columns: list[tuple], choices: dict | None = None,
              row_height: int | None = None):
    """columns: (key, header, width, editable). choices: key -> list of allowed values (drop-down)."""
    ws = wb.create_sheet(name)
    for j, (key, head, width, editable) in enumerate(columns, 1):
        c = ws.cell(row=1, column=j, value=head)
        c.font = HEAD_FONT
        c.fill = INPUT_FILL if editable else READ_FILL
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[get_column_letter(j)].width = width
    for i, row in enumerate(rows, 2):
        for j, (key, _, _, editable) in enumerate(columns, 1):
            v = row.get(key)
            if isinstance(v, (list, tuple, set)):
                v = "; ".join(map(str, v))
            c = ws.cell(row=i, column=j, value=v)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if editable:
                c.fill = INPUT_FILL
        if row_height:
            ws.row_dimensions[i].height = row_height
    ws.freeze_panes = "B2"
    for key, values in (choices or {}).items():
        j = next((k for k, c in enumerate(columns, 1) if c[0] == key), None)
        if j is None or not rows:
            continue
        # long lists go to a hidden sheet and are referenced by range (Excel limits inline lists to 255 chars)
        ref = list_range(wb, f"list_{key}", values)
        dv = DataValidation(type="list", formula1=ref, allow_blank=True, showErrorMessage=False)
        ws.add_data_validation(dv)
        col = get_column_letter(j)
        dv.add(f"{col}2:{col}{len(rows) + 1}")
    return ws


def list_range(wb: Workbook, name: str, values: list[str]) -> str:
    sheet = "lists"
    ws = wb[sheet] if sheet in wb.sheetnames else wb.create_sheet(sheet)
    ws.sheet_state = "hidden"
    col = ws.max_column + 1 if ws.max_row > 1 or ws["A1"].value else 1
    ws.cell(row=1, column=col, value=name)
    for i, v in enumerate(values, 2):
        ws.cell(row=i, column=col, value=v)
    L = get_column_letter(col)
    return f"'{sheet}'!${L}$2:${L}${len(values) + 1}"


def save(wb: Workbook, path: Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(f".{path.name}.tmp.xlsx")
    wb.save(tmp)
    tmp.replace(path)
    return path


def returned_files(folder: Path) -> list[Path]:
    folder = Path(folder)
    if not folder.is_dir():
        return []
    return sorted(f for f in folder.rglob("*.xlsx") if not f.name.startswith(("~$", ".")))


def read_sheet(path: Path, sheet: str) -> tuple[str | None, list[dict], dict]:
    """(reviewer_id, rows as dicts keyed by header, provenance). Header names are the visible headers."""
    wb = load_workbook(path, data_only=True, read_only=True)
    reviewer = None
    if "README" in wb.sheetnames:
        reviewer = wb["README"]["B3"].value
    reviewer = (str(reviewer).strip() if reviewer else None) or None
    prov = {"file": str(path), "sha256": file_sha256(path), "imported_at": now_iso(),
            "reviewer_from": "README!B3" if reviewer else "file name"}
    if reviewer is None:
        reviewer = Path(path).stem
    if sheet not in wb.sheetnames:
        return reviewer, [], {**prov, "error": f"no sheet {sheet!r}"}
    ws = wb[sheet]
    it = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(it, [])]
    rows = []
    for vals in it:
        if vals is None or all(v is None for v in vals):
            continue
        rows.append({h: v for h, v in zip(header, vals) if h})
    wb.close()
    return reviewer, rows, prov


def split_list(v) -> list[str]:
    if v is None:
        return []
    return [x.strip() for x in str(v).replace(",", ";").split(";") if x.strip()]
