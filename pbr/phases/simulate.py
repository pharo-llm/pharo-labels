"""A SIMULATED expert, only to test that completed workbooks are imported, validated and analysed correctly
before real experts answer. It is never a validation: the reviewer is named ``SIMULATED-not-a-human``, every
phase that reads it warns and stays provisional. Delete ``*/returned/SIMULATED*.xlsx`` before real use.

    ./run.sh simulate-expert validation     fills results/08_validation/expert/returned/SIMULATED.xlsx
    ./run.sh simulate-expert verification   fills results/12_verification/expert/returned/SIMULATED.xlsx

Validation simulation: for each item, the AI majority label with probability 0.7, otherwise another behaviour of
the same category (or any) — so human/AI metrics are neither perfect nor random. Relevance ratings random.
"""
from __future__ import annotations

import ast

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from .. import taxonomy as T
from ..common import BY_NUM, num_of, phase_dirname, results_dir

REVIEWER = "SIMULATED-not-a-human"


def main(cfg: dict, args: list[str]) -> int:
    arg = (args[0] if args else "validation")
    key = BY_NUM[int(arg)][1] if arg.isdigit() else arg
    rng = np.random.default_rng(int(cfg["project"]["seed"]) + 99)
    if key == "validation":
        d = results_dir(cfg) / phase_dirname("validation")
        src = d / "expert" / "expert_workbook.xlsx"
        if not src.exists():
            raise SystemExit("run the validation phase first (it writes expert/expert_workbook.xlsx)")
        tax = T.load(results_dir(cfg) / phase_dirname("taxonomy") / f"taxonomy_{cfg['taxonomy'].get('version', 'v1')}.json")
        items = pd.read_csv(d / "item_agreement.csv", converters={"majority": ast.literal_eval}).set_index("item_id")
        names = {b["id"]: f"{b['id']} {b['name']}" for b in T.active(tax, "behavior")}
        by_cat = {}
        for b in T.active(tax, "behavior"):
            by_cat.setdefault(b["category_id"], []).append(b["id"])
        wb = load_workbook(src)
        wb["README"]["B3"] = REVIEWER
        ws = wb["Method Annotation"]
        head = [c.value for c in ws[1]]
        col = {h: i + 1 for i, h in enumerate(head)}
        for r in range(2, ws.max_row + 1):
            key = ws.cell(r, col["row_key"]).value
            maj = list(items.loc[key, "majority"]) if key in items.index else []
            maj = [m for m in maj if m in names] or [rng.choice(list(names))]
            if rng.random() > 0.7:
                cat = T.behavior_path(tax, maj[0])[0]
                pool = [b for b in by_cat.get(cat, []) if b != maj[0]] or list(names)
                maj = [rng.choice(pool)] + maj[1:2]
            for i, b in enumerate(maj[:3], 1):
                ws.cell(r, col[f"behavior_{i}"]).value = names[b]
            ws.cell(r, col["confidence"]).value = rng.choice(["low", "medium", "high"])
            ws.cell(r, col["notes"]).value = "simulated"
        ws = wb["Behavior Relevance"]
        head = [c.value for c in ws[1]]
        col = {h: i + 1 for i, h in enumerate(head)}
        for r in range(2, ws.max_row + 1):
            ws.cell(r, col["relevance_pharo"]).value = str(int(rng.integers(1, 6)))
            ws.cell(r, col["relevance_benchmark"]).value = str(int(rng.integers(1, 6)))
        out = d / "expert" / "returned" / "SIMULATED.xlsx"
    elif key == "verification":
        d = results_dir(cfg) / phase_dirname("verification")
        src = d / "expert" / "benchmark_review.xlsx"
        if not src.exists():
            raise SystemExit("run the verification phase first (it writes expert/benchmark_review.xlsx)")
        wb = load_workbook(src)
        wb["README"]["B3"] = REVIEWER
        ws = wb["Benchmark Task Validation"]
        head = [c.value for c in ws[1]]
        col = {h: i + 1 for i, h in enumerate(head)}
        for r in range(2, ws.max_row + 1):
            for k in ("clarity", "correctness", "realism", "test_quality"):
                ws.cell(r, col[k]).value = str(int(rng.integers(3, 6)))
            ws.cell(r, col["decision"]).value = "accept" if rng.random() < 0.85 else "revise"
        out = d / "expert" / "returned" / "SIMULATED.xlsx"
    else:
        raise SystemExit("simulate-expert validation | verification")
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"wrote {out} (reviewer {REVIEWER}). Rerun ./run.sh {key} to import it; delete it before real use.")
    return 0
